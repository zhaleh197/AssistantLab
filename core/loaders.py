"""استخراج متن از انواع منابع.

خروجی هر تابع لیستی از (location, text) است؛ location مثل «صفحه ۳» برای ارجاع استفاده می‌شود.
"""
import importlib.util
import io
import ipaddress
import socket
import subprocess
import sys
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup

import config
from core.text import clean

Segments = list[tuple[str, str]]

TEXT_EXTS = {".txt", ".md", ".markdown", ".json", ".xml", ".yaml", ".yml", ".log", ".rst"}
SUPPORTED_EXTS = sorted(TEXT_EXTS | {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".html", ".htm"})


class UnsupportedSource(ValueError):
    pass


def decode_bytes(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1256"):
        try:
            text = data.decode(enc)
            if enc != "utf-16" or "\x00" not in text:
                return text
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> Segments:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    segs = [(f"صفحه {i}", page.extract_text() or "") for i, page in enumerate(reader.pages, 1)]
    segs = [(loc, t) for loc, t in segs if t.strip()]
    if not segs:
        raise UnsupportedSource("از این PDF متنی استخراج نشد؛ احتمالاً اسکن‌شده است (OCR هنوز پشتیبانی نمی‌شود).")
    return segs


def _docx(data: bytes) -> Segments:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return [("", "\n\n".join(parts))]


def _pptx(data: bytes) -> Segments:
    from pptx import Presentation

    segs = []
    for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        if texts:
            segs.append((f"اسلاید {i}", "\n".join(texts)))
    return segs


def _table_rows_to_text(df) -> str:
    """هر ردیف جدول به شکل «ستون: مقدار؛ ...» تا مدل معنای هر خانه را بفهمد."""
    df = df.dropna(how="all").fillna("")
    lines = []
    for _, row in df.iterrows():
        cells = [f"{col}: {val}" for col, val in row.items() if str(val).strip()]
        if cells:
            lines.append("؛ ".join(cells))
    return "\n\n".join(lines)


def _xlsx(data: bytes) -> Segments:
    import pandas as pd

    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=str)
    return [(f"شیت {name}", _table_rows_to_text(df)) for name, df in sheets.items()]


def _csv(data: bytes) -> Segments:
    import pandas as pd

    return [("", _table_rows_to_text(pd.read_csv(io.StringIO(decode_bytes(data)), dtype=str)))]


@dataclass
class HtmlDoc:
    title: str
    description: str
    text: str
    links: list[str]

    @property
    def full_text(self) -> str:
        return f"{self.description}\n\n{self.text}".strip()


def _html_to_text(html: str | bytes, base_url: str = "") -> HtmlDoc:
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else ""
    meta = soup.find("meta", attrs={"name": "description"}) or soup.find("meta", attrs={"property": "og:description"})
    description = (meta.get("content") or "").strip() if meta else ""
    # لینک‌ها قبل از حذف منو و فوتر جمع می‌شوند (بیشتر لینک‌های داخلی همان‌جا هستند)
    links = [urljoin(base_url, a["href"]) for a in soup.find_all("a", href=True)] if base_url else []
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "aside", "form", "svg", "template"]):
        tag.decompose()
    body = soup.body or soup
    main = soup.find("main") or soup.find("article")
    main_text = main.get_text(separator="\n") if main else ""
    # اگر <main> تقریباً خالی باشد، کل body استفاده می‌شود
    text = main_text if len(main_text.strip()) >= 200 else body.get_text(separator="\n")
    return HtmlDoc(title, description, text, links)


def extract_file(filename: str, data: bytes) -> Segments:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        segs = _pdf(data)
    elif ext == ".docx":
        segs = _docx(data)
    elif ext == ".pptx":
        segs = _pptx(data)
    elif ext == ".xlsx":
        segs = _xlsx(data)
    elif ext == ".csv":
        segs = _csv(data)
    elif ext in (".html", ".htm"):
        segs = [("", _html_to_text(data).full_text)]
    elif ext in TEXT_EXTS:
        segs = [("", decode_bytes(data))]
    elif ext in (".doc", ".ppt", ".xls"):
        raise UnsupportedSource(f"فرمت قدیمی {ext} پشتیبانی نمی‌شود؛ فایل را با فرمت جدید (مثلاً {ext}x) ذخیره کنید.")
    else:
        raise UnsupportedSource(f"فرمت {ext or 'بدون پسوند'} پشتیبانی نمی‌شود.")
    return [(loc, t) for loc, t in segs if t.strip()]


# ---------- لینک وب ----------

@lru_cache(maxsize=256)
def is_public_host(hostname: str | None) -> bool:
    """فقط آدرس‌های عمومی اینترنت مجازند (جلوگیری از SSRF در سیستم چندسازمانی)."""
    if not hostname:
        return False
    try:
        infos = socket.getaddrinfo(hostname, None)
    except socket.gaierror:
        return False
    return all(ipaddress.ip_address(info[4][0]).is_global for info in infos)


def _assert_public_host(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise UnsupportedSource("آدرس باید با http:// یا https:// شروع شود.")
    try:
        socket.getaddrinfo(parsed.hostname, None)
    except socket.gaierror:
        raise UnsupportedSource("آدرس سایت پیدا نشد.")
    if not is_public_host(parsed.hostname):
        raise UnsupportedSource("دسترسی به آدرس‌های داخلی شبکه مجاز نیست.")


JS_SITE_MSG = (
    "این صفحه محتوایش را با جاوااسکریپت می‌سازد و برای خواندنش مرورگر Playwright لازم است. "
    "نصب: pip install playwright و سپس python -m playwright install chromium-headless-shell"
)


def _render_with_browser(url: str) -> bytes:
    """اجرای صفحه در مرورگر headless (در پروسه‌ی جدا) و برگرداندن HTML نهایی."""
    if importlib.util.find_spec("playwright") is None:
        raise UnsupportedSource(JS_SITE_MSG)
    try:
        proc = subprocess.run(
            [sys.executable, "-m", "core.render_page", url],
            cwd=config.ROOT, capture_output=True, timeout=config.RENDER_TIMEOUT_MS / 1000 + 30,
        )
    except subprocess.TimeoutExpired:
        raise UnsupportedSource("بارگذاری صفحه در مرورگر بیش از حد طول کشید.")
    if proc.returncode != 0:
        err = proc.stderr.decode("utf-8", errors="replace")
        if "Executable doesn't exist" in err:
            raise UnsupportedSource(JS_SITE_MSG)
        raise UnsupportedSource(f"باز کردن صفحه با مرورگر ناموفق بود: {err.strip().splitlines()[-1][:200] if err.strip() else ''}")
    if len(proc.stdout) > config.MAX_URL_BYTES:
        raise UnsupportedSource("حجم این صفحه بیش از حد مجاز است.")
    return proc.stdout


USER_AGENT = "Mozilla/5.0 (compatible; AssistantLab/1.0; document fetcher)"


@dataclass
class WebPage:
    url: str              # آدرس نهایی (بعد از ریدایرکت)
    title: str
    segments: Segments
    links: list[str] = field(default_factory=list)


def _http_get(url: str) -> tuple[str, str, bytes]:
    """(آدرس نهایی، Content-Type، بدنه) را برمی‌گرداند."""
    # ریدایرکت‌ها دستی دنبال می‌شوند تا مقصد هر مرحله هم بررسی شود
    for _ in range(5):
        _assert_public_host(url)
        resp = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=20, stream=True, allow_redirects=False)
        if resp.is_redirect:
            url = urljoin(url, resp.headers["Location"])
            continue
        break
    else:
        raise UnsupportedSource("تعداد ریدایرکت‌ها زیاد است.")
    resp.raise_for_status()

    data = b""
    for block in resp.iter_content(64 * 1024):
        data += block
        if len(data) > config.MAX_URL_BYTES:
            raise UnsupportedSource("حجم این صفحه بیش از حد مجاز است.")
    return url, resp.headers.get("Content-Type", "").lower(), data


def load_page(url: str) -> WebPage:
    url, ctype, data = _http_get(url.strip())
    path_name = Path(urlparse(url).path).name or url
    if "pdf" in ctype or url.lower().endswith(".pdf"):
        return WebPage(url, path_name, _pdf(data))
    if "html" in ctype or not ctype:
        # بایت خام به BeautifulSoup داده می‌شود تا کدگذاری را از meta صفحه تشخیص دهد
        doc = _html_to_text(data, url)
        if len(clean(doc.text)) < config.MIN_STATIC_TEXT_CHARS:
            # احتمالاً سایت تک‌صفحه‌ای (SPA) است و محتوا با جاوااسکریپت ساخته می‌شود
            try:
                rendered = _html_to_text(_render_with_browser(url), url)
            except UnsupportedSource:
                if not clean(doc.text):
                    raise
            else:
                if len(clean(rendered.text)) > len(clean(doc.text)):
                    rendered.title = rendered.title or doc.title
                    rendered.description = rendered.description or doc.description
                    doc = rendered
        return WebPage(url, doc.title or url, [("", doc.full_text)], doc.links)
    return WebPage(url, path_name, extract_file(path_name if Path(path_name).suffix else path_name + ".txt", data))


def extract_url(url: str) -> tuple[str, Segments]:
    """(عنوان، بخش‌ها) را برای یک صفحه برمی‌گرداند."""
    page = load_page(url)
    return page.title, page.segments
